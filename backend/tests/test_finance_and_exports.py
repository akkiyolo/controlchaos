"""Tests for Finance Views (Variance Studio, Treasury, Reconciliation) and Exports (Excel, M-queries, Audit Pack)."""

from __future__ import annotations

import io

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.entities import (
    Account,
    Budget,
    Dataset,
    GLEntry,
    Run,
    Score,
    SubledgerEntry,
    TreasuryPosition,
)


def test_finance_variance_studio(client: TestClient, db_session: Session, auth_headers: dict):
    # Setup test dataset and COA
    coa = Account(code="5100", name="Software & Subscriptions", type="expense", normal_balance="debit")
    db_session.add(coa)

    ds = Dataset(
        id="ds-fin-test",
        name="Finance Test Dataset",
        seed=42,
        period_start="2025-01",
        period_end="2025-01",
        created_by="admin@controlchaos.io",
        is_baseline=True,
    )
    db_session.add(ds)

    gl = GLEntry(
        dataset_id="ds-fin-test",
        entry_id="GL-101",
        posting_date="2025-01-15",
        period="2025-01",
        entity_id="ENT-01",
        account_code="5100",
        debit=15000.0,
        credit=0.0,
        currency="USD",
        description="AWS Hosting Monthly",
        source_system="ERP",
        batch_id="B-1",
        created_by="system",
        approved_by="manager",
    )
    db_session.add(gl)

    b = Budget(
        dataset_id="ds-fin-test",
        period="2025-01",
        entity_id="ENT-01",
        account_code="5100",
        amount=10000.0,
    )
    db_session.add(b)
    db_session.commit()

    # Query Variance Studio
    res = client.get("/api/v1/finance/variance-studio?dataset_id=ds-fin-test&period=2025-01", headers=auth_headers["admin"])
    assert res.status_code == 200
    data = res.json()
    assert data["dataset_id"] == "ds-fin-test"
    assert len(data["rows"]) >= 1
    target_row = next(r for r in data["rows"] if r["account_code"] == "5100")
    assert target_row["actual"] == 15000.0
    assert target_row["budget"] == 10000.0
    assert target_row["variance_amount"] == 5000.0
    assert target_row["is_adverse"] is True

    # Test AI commentary endpoint
    ai_res = client.post(
        "/api/v1/finance/variance-studio/ai-commentary",
        headers=auth_headers["admin"],
        json={
            "dataset_id": "ds-fin-test",
            "account_code": "5100",
            "period": "2025-01",
            "actual": 15000.0,
            "budget": 10000.0,
        },
    )
    assert ai_res.status_code == 200
    ai_data = ai_res.json()
    assert ai_data["account_code"] == "5100"
    assert "management_commentary" in ai_data


def test_finance_treasury_and_recon(client: TestClient, db_session: Session, auth_headers: dict):
    # Treasury positions
    tp = TreasuryPosition(
        dataset_id="ds-fin-test",
        as_of_date="2025-01-31",
        entity_id="ENT-01",
        currency="USD",
        cash_balance=10000000.0,
        hqla=8000000.0,
        outflows_30d=5000000.0,
        inflows_30d=2000000.0,
        funding_source="Commercial Paper",
        funding_amount=3000000.0,
        maturity_bucket="30d",
    )
    db_session.add(tp)

    # Subledger entry
    sub = SubledgerEntry(
        dataset_id="ds-fin-test",
        ref_id="INV-999",
        gl_entry_id=None,
        subledger_type="AR",
        counterparty="Acme Corp",
        amount=25000.0,
        currency="USD",
        value_date="2025-01-10",
        status="open",
    )
    db_session.add(sub)
    db_session.commit()

    # Treasury endpoint
    tres = client.get("/api/v1/finance/treasury?dataset_id=ds-fin-test", headers=auth_headers["admin"])
    assert tres.status_code == 200
    tdata = tres.json()
    assert tdata["total_cash_usd"] >= 10000000.0
    assert tdata["lcr_ratio"] > 0

    # Recon endpoint
    rres = client.get("/api/v1/finance/reconciliation?dataset_id=ds-fin-test", headers=auth_headers["admin"])
    assert rres.status_code == 200
    rdata = rres.json()
    assert rdata["total_breaks"] >= 1
    assert len(rdata["subledgers"]) == 5

    # Resolve break
    resolve_res = client.post(
        "/api/v1/finance/reconciliation/resolve-break",
        headers=auth_headers["admin"],
        json={
            "break_id": "BRK-INV-999",
            "action": "investigate",
            "resolution_notes": "Pending counterparty payment trace confirmation.",
        },
    )
    assert resolve_res.status_code == 200


def test_exports_excel_and_mqueries(client: TestClient, db_session: Session, auth_headers: dict):
    run = Run(
        id="run-export-test",
        dataset_id="ds-fin-test",
        control_set_snapshot={},
        status="completed",
        config={},
    )
    db_session.add(run)

    score = Score(
        run_id="run-export-test",
        overall_detection_rate=0.85,
        weighted_detection_rate=0.92,
        precision=0.88,
        recall=0.85,
        f1=0.86,
        cost_of_misses=4500.0,
        false_positive_rate=0.04,
        by_class={"amount_mismatch": {"total": 5, "caught": 4, "rate": 0.8}},
        by_control={},
    )
    db_session.add(score)
    db_session.commit()

    # Excel export
    excel_res = client.get("/api/v1/exports/excel?run_id=run-export-test", headers=auth_headers["admin"])
    assert excel_res.status_code == 200
    assert excel_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # Verify openpyxl can load the exported bytes
    wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
    assert "Executive Summary" in wb.sheetnames
    assert "Blind Spot Heatmap" in wb.sheetnames

    # Power BI M-queries
    pbi_res = client.get("/api/v1/exports/powerbi/m-queries", headers=auth_headers["admin"])
    assert pbi_res.status_code == 200
    pbi_data = pbi_res.json()
    assert "m_queries" in pbi_data
    assert "Variance_Studio" in pbi_data["m_queries"]

    # Audit Pack
    audit_res = client.get("/api/v1/exports/audit-pack?run_id=run-export-test", headers=auth_headers["admin"])
    assert audit_res.status_code == 200
    audit_data = audit_res.json()
    assert audit_data["chain_verification"]["algorithm"] == "SHA-256"
